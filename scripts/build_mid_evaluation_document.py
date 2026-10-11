from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs_v2" / "mid_evaluation"
DOCX = OUT / "principles_and_practices.docx"

BLACK = "000000"
NAVY = "17364A"
TEAL = "397283"
PALE = "EDF4F6"
GRID = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def border_table(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRID)
        borders.append(element)
    tbl_pr.append(borders)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")


def mark_header_row(row):
    row_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_properties.append(header)


def style_run(run, size=9.4, bold=False, color=BLACK, font="Aptos"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)


def add_body(doc, text, bold_lead=None, size=9.4, after=4):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.06
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        style_run(lead, size=size, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        style_run(rest, size=size)
    else:
        run = paragraph.add_run(text)
        style_run(run, size=size)
    return paragraph


def add_bullet(doc, text, level=0, size=9.1):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.02
    style_run(paragraph.add_run(text), size=size)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    style_run(run, size=13.2 if level == 1 else 10.5, bold=True)
    return paragraph


def add_table(doc, headers, rows, widths, font_size=8.1):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    border_table(table)
    mark_header_row(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, NAVY)
        set_width(cell, widths[index])
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_run(paragraph.add_run(header), size=font_size, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            if row_index % 2:
                shade(cells[column], PALE)
            set_width(cells[column], widths[column])
            set_cell_margins(cells[column])
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cells[column].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            style_run(paragraph.add_run(value), size=font_size, bold=(column == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_footer(section):
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(7.25))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mark_header_row(table.rows[0])
    left, right = table.rows[0].cells
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(left.paragraphs[0].add_run("CS5013  IITM Explainable Slotbook"), size=8, color="5D6F79")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    right.paragraphs[0]._p.append(field)


def suppress_paragraph_border(paragraph):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        paragraph_properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    borders.append(bottom)


OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
add_footer(section)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.4)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(23)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(4)
suppress_paragraph_border(title)
style_run(title.add_run("Engineering Principles and Practices"), size=23, bold=True, font="Aptos Display")
meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(8)
style_run(meta.add_run("IITM Explainable Slotbook  |  CS5013 Mid Evaluation  |  11 October 2026"), size=9, bold=True, color=TEAL)
style_run(meta.add_run("\nNishanth Senthil Kumar EE23B049  |  Sparsh Gupta EE23B117"), size=9, color="526773")

add_body(doc, "This project applies the course principles as an auditable engineering loop: reviewed prompt, focused implementation, automated verification, documentation, and a coherent Git commit. The system collects synthetic course preferences, maps IITM constraints into UniTime, checks results against CP-SAT, and publishes deterministic evidence for every slot assignment.", size=9.5, after=5)

add_heading(doc, "1 Prompt Engineering")
add_body(doc, "Prompts are versioned engineering inputs, not disposable chat messages. Each substantial task is recorded in prompts/PROMPT_LOG.md with CoSTAR structure, negative constraints, review notes, acceptance checks, and the commit it influenced.")
add_table(doc, ["CoSTAR field", "How we apply it"], [
    ("Context", "Name the current architecture, data provenance, relevant files, and known constraints."),
    ("Objective", "State one observable outcome, such as a valid fixture or a passing API contract."),
    ("Style", "Require small cohesive modules, stable schemas, and maintainable test assertions."),
    ("Tone", "Use precise claims; label synthetic data, uncertainty, and heuristic results honestly."),
    ("Audience", "Target the administrator, evaluator, student, professor, or future maintainer."),
    ("Response", "Specify files, commands, tests, documentation, and acceptance evidence."),
], [1.35, 5.75], font_size=8.2)
add_body(doc, "Review gate. We check specificity, hidden assumptions, privacy, architecture fit, failure modes, and verifiability before accepting generated work. Prompt revisions and rejected outputs remain in the log so Git records the reasoning, not only the final diff.", bold_lead="Review gate.")
add_body(doc, "Representative negative prompts. Do not use real student identities; do not hand-author the final timetable; do not let an LLM decide or invent assignment reasons; do not change CSV contracts silently; do not claim UniTime is globally optimal; and do not publish until validation and CI succeed.", bold_lead="Representative negative prompts.")

doc.add_page_break()
add_heading(doc, "2 Structured Debugging")
add_body(doc, "We use reproduce, observe, hypothesize, isolate, correct, test, regress, and document. A failure is first reduced to one command or fixture. Competing hypotheses are written before editing, and the smallest correction is made in the module that owns the behaviour.")
add_table(doc, ["Case", "Evidence and correction"], [
    ("Objective mismatch", "UniTime reported 12 for a schedule whose domain cost was 20. Inspection showed cpsolver scores time preferences relative to each class minimum. We restored the constant offset and added an independent penalty-equivalence regression test."),
    ("Non deterministic tests", "Timetables differed between JVM runs. A randomised map iteration order changed constraint insertion order. Stable pair ordering fixed it; the same iteration cap now reproduces the same timetable."),
    ("Apps Script 404", "The Vercel bridge followed a POST ContentService redirect automatically and received HTML. We validated the script.googleusercontent.com redirect and retrieved it with GET; a Node contract test now pins that behaviour."),
    ("Weak showcase", "The first demo had zero student clashes, so it failed the presentation goal despite being valid. Six synthetic pair coverage profiles made all 28 elective pairs positive, forcing one explainable minimum weight overlap."),
], [1.55, 5.55], font_size=7.9)

add_heading(doc, "3 Version Control")
add_body(doc, "Commits are coherent verified increments. Prompts are committed with the implementation they influence; secrets, IDE state, build directories, and temporary exports are excluded. Before every commit we inspect the diff and run the nearest tests, then the applicable full suite.")
add_table(doc, ["Commit", "Purpose and verification boundary"], [
    ("487a534", "Added the showcase dataset, workbook, validator, CI generation, and reviewed prompt."),
    ("f4330ee", "Rejected the weak scenario and forced a genuine explainable optimisation trade off."),
    ("8269d7f", "Retained solver generated outputs only after UniTime matched CP-SAT objective 9."),
], [1.35, 5.75], font_size=8.1)
add_body(doc, "Branch protection is supported by GitHub Actions on every push and pull request. The workflow gives future refactors the same executable acceptance criteria as the original implementation.")

doc.add_page_break()
add_heading(doc, "4 Testing Refactoring and Documentation")
add_body(doc, "Tests target observable guarantees instead of one fragile timetable. Multiple schedules can have equal cost, so assertions check completeness, domains, fixed slots, real time overlaps, labs, rooms, protected pairs, objective recomputation, explanation coverage, and failure codes.")
add_table(doc, ["Evidence", "Verified result"], [
    ("Java model", "52 JUnit tests across parsing, validation, both solvers, exports, generated scenarios, independent cost checks, and brute force small cases."),
    ("Portal", "7 Node API contract tests plus a static validator covering 18 catalog courses, 17 slots, required assets, and CSV headers."),
    ("Showcase", "11 courses, 36 synthetic students, 144 ranked rows, 6 protected pairs, and all 28 elective co demand pairs covered."),
    ("Cross solver check", "UniTime produced objective 9 in 10.108 seconds; CP-SAT proved objective 9 optimal in 0.199 seconds."),
    ("Continuous integration", "CI compiles Java 17, runs both test suites and validators, executes both solvers, and uploads the generated slotbooks."),
], [1.55, 5.55], font_size=7.9)
add_body(doc, "Refactoring. TimeGrid is the single definition of overlap; Objective is the single definition of cost; TimetableSolver isolates engines; UniTime imports stay in solver.unitime; and CLI, JavaFX, and the portal depend on stable outputs rather than solver variables. The portal similarly separates static forms, server-side API validation, Apps Script persistence, and publication.", bold_lead="Refactoring.")
add_body(doc, "Documentation. The repository includes a man-page-style manual with name, synopsis, options, inputs, outputs, exit status, examples, and limitations; an architecture flow; code walkthrough; test plan; verification record; debugging guide; file map; prompt register; and this evaluation runbook.", bold_lead="Documentation.")

add_heading(doc, "Appendix Sample Reviewed Prompts")
add_body(doc, "Showcase prompt. Context: the portal exports solver-compatible preferences. Objective: create a compact synthetic case that demonstrates every important constraint and one genuine trade off. Style: evaluator-facing CSVs and a workbook. Audience: CS5013 evaluators. Response: fixture, validator, both solver outputs, CI artifact, and public explanation. Negative constraints: no real IDs, no hand-authored result, no hidden hard-rule violations, and no optimality claim without CP-SAT proof.", bold_lead="Showcase prompt.", size=8.7, after=3)
add_body(doc, "Live change prompt. Context: add a positive-class-size invariant without changing solver behaviour. Objective: reject class_size less than one with a stable issue code. Style: smallest change in InputValidator plus one focused JUnit case. Audience: maintainers and evaluator. Response: code diff, focused command, full-suite command, and documentation impact. Negative constraints: do not weaken existing validation, edit the CSV parser, or assert an error-message sentence instead of the issue code.", bold_lead="Live change prompt.", size=8.7, after=2)

doc.core_properties.title = "Engineering Principles and Practices"
doc.core_properties.subject = "CS5013 Mid Evaluation"
doc.core_properties.author = "Nishanth Senthil Kumar and Sparsh Gupta"
doc.core_properties.keywords = "CS5013, prompt engineering, debugging, Git, testing, timetabling"
doc.save(DOCX)
print(DOCX)
